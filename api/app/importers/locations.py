"""Spain locations import: INE register + INE population + GeoNames (coordinates, names en/uk/ru).

Pipeline (spec §3):
  1. INE municipality register -> communities, provinces, ~8 100 municipalities with codes.
  2. INE padrón -> population.
  3. GeoNames ES dump: for Spain `admin3` is the 5-digit INE municipality code, so each municipality is
     joined by code. Coordinates come from the populated place (town centre), not the ADM3 centroid.
  4. GeoNames alternateNames -> en/uk/ru names and co-official/Castilian aliases.
  5. seeds/locations_overrides.json -> Castilian exonyms and manually checked uk/ru names for the top cities.
  6. Slugs: from the Spanish name, globally unique; on collision the smaller municipality gets the province
     suffix (villanueva-toledo). Provinces/communities are prefixed (provincia-valencia, comunidad-madrid).
     Existing slugs are never changed by a re-import.
"""

import io
import json
import re
import unicodedata
import urllib.request
import zipfile
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from geoalchemy2 import WKTElement
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.slug import slugify, transliterate
from app.models import Category, Location, Section

ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data" / "raw"
REGIONS_FILE = ROOT / "seeds" / "regions.json"
OVERRIDES_FILE = ROOT / "seeds" / "locations_overrides.json"

SOURCES = {
    "ine_register": ("diccionario25.xlsx", "https://www.ine.es/daco/daco42/codmun/diccionario25.xlsx"),
    "ine_population": ("pobmun22.xlsx", "https://www.ine.es/pob_xls/pobmun22.xlsx"),
    "geonames": ("ES.zip", "https://download.geonames.org/export/dump/ES.zip"),
    "geonames_alt": (
        "alternateNamesV2.zip",
        "https://download.geonames.org/export/dump/alternateNamesV2.zip",
    ),
}

NAME_LANGS = ("en", "uk", "ru")
# Latin-script languages whose names become search aliases (Castilian + co-official).
ALIAS_LANGS = ("es", "ca", "gl", "eu", "oc", "ast", "an")
ARTICLES = {"a", "o", "as", "os", "el", "la", "los", "las", "les", "els", "es", "sa", "ses", "lo", "l'"}
PLACE_RANK = {"PPLC": 0, "PPLA": 1, "PPLA2": 2, "PPLA3": 3, "PPLA4": 4, "PPL": 5}
MAX_ALIASES = 25
# GeoNames marks some facilities as populated places. Industrial estates and urbanizaciones stay:
# they are real workplaces / neighbourhoods (spec §3: a listing may point to a polígono).
NOT_A_LOCALITY = re.compile(
    r"\b(polideportivo|estadio|aeropuerto|aeroport|camping|hotel|hospital|cementerio|estacion)\b"
)

Log = Callable[[str], None]


# ---------------------------------------------------------------------------- helpers


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def fix_article(part: str) -> str:
    """INE puts articles last: 'Coruña, A' -> 'A Coruña', "Hospitalet, L'" -> "L'Hospitalet"."""
    match = re.match(r"^(.*), ([^,]+)$", part.strip())
    if not match or match.group(2).lower() not in ARTICLES:
        return part.strip()
    base, article = match.groups()
    article = article[0].upper() + article[1:]
    return f"{article}{base}" if article.lower() == "l'" else f"{article} {base}"


def ensure_sources(raw_dir: Path, download: bool, log: Log) -> dict[str, Path]:
    raw_dir.mkdir(parents=True, exist_ok=True)
    paths = {}
    for key, (filename, url) in SOURCES.items():
        path = raw_dir / filename
        if not path.exists():
            if not download:
                raise FileNotFoundError(f"{path} is missing; run with --download or fetch {url}")
            log(f"downloading {url}")
            urllib.request.urlretrieve(url, path)  # noqa: S310 - fixed https URLs
        paths[key] = path
    return paths


# ---------------------------------------------------------------------------- sources


@dataclass
class IneMunicipality:
    code: str  # CPRO + CMUN
    community: str  # CODAUTO
    province: str  # CPRO
    raw_name: str
    parts: list[str]


def read_ine_register(path: Path) -> list[IneMunicipality]:
    sheet = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    result = []
    for row in sheet.iter_rows(values_only=True):
        codauto, cpro, cmun, _dc, nombre = (row + (None,) * 5)[:5]
        if not (codauto and cpro and cmun and nombre) or not str(cmun).isdigit():
            continue
        name = str(nombre).strip()
        result.append(
            IneMunicipality(
                code=f"{int(cpro):02d}{int(cmun):03d}",
                community=f"{int(codauto):02d}",
                province=f"{int(cpro):02d}",
                raw_name=name,
                parts=[fix_article(p) for p in name.split("/") if p.strip()],
            )
        )
    return result


def read_ine_population(path: Path) -> dict[str, int]:
    sheet = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    population = {}
    for row in sheet.iter_rows(values_only=True):
        cpro, _prov, cmun, _name, total = (row + (None,) * 5)[:5]
        if cpro and cmun and str(cmun).isdigit() and isinstance(total, int | float):
            population[f"{int(cpro):02d}{int(cmun):03d}"] = int(total)
    return population


@dataclass
class GeoFeature:
    id: str
    name: str
    lat: float
    lon: float
    fcode: str
    population: int


@dataclass
class GeoNames:
    adm3: dict[str, GeoFeature]  # INE municipality code -> ADM3
    places: dict[str, list[GeoFeature]]  # INE municipality code -> populated places
    # Fallback for municipalities created after GeoNames was updated (no INE code there yet):
    places_by_name: dict[tuple[str, str], list[GeoFeature]]  # (GeoNames admin2, normalized name)
    province_admin2: dict[str, str]  # INE CPRO -> GeoNames admin2
    by_id: dict[str, GeoFeature] = field(default_factory=dict)  # populated places, for overrides


def read_geonames(path: Path) -> GeoNames:
    adm3: dict[str, GeoFeature] = {}
    places: dict[str, list[GeoFeature]] = defaultdict(list)
    by_name: dict[tuple[str, str], list[GeoFeature]] = defaultdict(list)
    by_id: dict[str, GeoFeature] = {}
    votes: dict[str, Counter[str]] = defaultdict(Counter)
    with zipfile.ZipFile(path) as z, z.open("ES.txt") as fh:
        for line in io.TextIOWrapper(fh, encoding="utf-8"):
            c = line.rstrip("\n").split("\t")
            is_place = c[6] == "P" and c[7] in PLACE_RANK
            if not (is_place or c[7] == "ADM3"):
                continue
            feature = GeoFeature(c[0], c[1], float(c[4]), float(c[5]), c[7], int(c[14] or 0))
            code, admin2 = c[12], c[11]
            if is_place:
                by_name[(admin2, norm(c[1]))].append(feature)
                by_id[feature.id] = feature
            if len(code) != 5 or not code.isdigit():
                continue
            votes[code[:2]][admin2] += 1
            if c[7] == "ADM3":
                adm3[code] = feature
            else:
                places[code].append(feature)
    province_admin2 = {cpro: counter.most_common(1)[0][0] for cpro, counter in votes.items()}
    return GeoNames(adm3, places, by_name, province_admin2, by_id)


def find_place_by_name(m: IneMunicipality, geo: GeoNames) -> GeoFeature | None:
    admin2 = geo.province_admin2.get(m.province)
    if admin2 is None:
        return None
    for part in m.parts:
        key = norm(part)
        without_article = re.sub(r"^(el|la|los|las|l|els|les|a|o|as|os|es|sa|ses) ", "", key)
        for variant in dict.fromkeys((key, without_article)):
            if candidates := geo.places_by_name.get((admin2, variant)):
                return min(candidates, key=lambda f: (PLACE_RANK[f.fcode], -f.population))
    return None


def read_alternate_names(
    path: Path, ids: set[str], log: Log
) -> dict[str, dict[str, list[tuple[str, bool, bool]]]]:
    """geonameid -> lang -> [(name, preferred, short)], without historic/colloquial names."""
    wanted_langs = set(NAME_LANGS) | set(ALIAS_LANGS)
    result: dict[str, dict[str, list[tuple[str, bool, bool]]]] = defaultdict(lambda: defaultdict(list))
    log("scanning GeoNames alternate names (takes a minute)")
    with zipfile.ZipFile(path) as z, z.open("alternateNamesV2.txt") as fh:
        for line in io.TextIOWrapper(fh, encoding="utf-8"):
            c = line.rstrip("\n").split("\t")
            if len(c) < 8 or c[1] not in ids or c[2] not in wanted_langs:
                continue
            if c[6] == "1" or c[7] == "1":  # colloquial / historic
                continue
            result[c[1]][c[2]].append((c[3], c[4] == "1", c[5] == "1"))
    return result


def best_name(candidates: Iterable[tuple[str, bool, bool]]) -> str | None:
    ranked = sorted(candidates, key=lambda n: (not n[1], n[2]))
    return ranked[0][0] if ranked else None


# ---------------------------------------------------------------------------- build


@dataclass
class Row:
    level: str
    ine_code: str | None
    parent_code: str | None  # INE code of the parent
    slug: str
    names: dict[str, str]
    aliases: list[str]
    population: int | None
    lon: float | None
    lat: float | None
    geonames_id: int | None = None


@dataclass
class ImportReport:
    comunidades: int = 0
    provincias: int = 0
    municipios: int = 0
    localidades: int = 0
    coords_from_place: int = 0
    coords_from_adm3: int = 0
    matched_by_name: list[str] = field(default_factory=list)
    approximate_coords: list[str] = field(default_factory=list)
    without_geonames: list[str] = field(default_factory=list)
    slug_collisions: list[str] = field(default_factory=list)
    override_warnings: list[str] = field(default_factory=list)
    top: list[Row] = field(default_factory=list)


def pick_place(m: IneMunicipality, candidates: list[GeoFeature]) -> GeoFeature | None:
    names = {norm(p) for p in m.parts}
    return min(
        candidates,
        key=lambda f: (norm(f.name) not in names, PLACE_RANK[f.fcode], -f.population),
        default=None,
    )


def clean_aliases(values: Iterable[str], exclude: str) -> list[str]:
    seen = {norm(exclude)}
    result = []
    for value in values:
        value = value.strip()
        key = norm(value)
        if value and key and key not in seen and not value.isdigit():
            seen.add(key)
            result.append(value)
    return result[:MAX_ALIASES]


def build_municipalities(
    register: list[IneMunicipality],
    population: dict[str, int],
    geo: GeoNames,
    alt_names: dict[str, dict[str, list[tuple[str, bool, bool]]]],
    overrides: dict[str, Any],
    reserved_slugs: set[str],
    province_slugs: dict[str, str],
    existing_slugs: dict[tuple[str, str], str],
    report: ImportReport,
) -> list[Row]:
    rows: list[Row] = []
    borrowed: list[tuple[Row, IneMunicipality, str]] = []  # rows that take coordinates from a neighbour
    for m in register:
        override = overrides.get(m.code, {})
        if (expect := override.get("expect")) and norm(expect) not in norm(m.raw_name):
            report.override_warnings.append(f"{m.code}: expected '{expect}', INE has '{m.raw_name}'")
            override = {}

        place = pick_place(m, geo.places.get(m.code, []))
        area = geo.adm3.get(m.code)
        if gid := override.get("geonames_id"):
            place = geo.by_id.get(str(gid))
            if place is None:
                report.override_warnings.append(f"{m.code}: GeoNames place {gid} not found")
        if place is None and area is None and (place := find_place_by_name(m, geo)):
            report.matched_by_name.append(f"{m.code} {m.raw_name} -> {place.name} ({place.id})")
        feature = place or area
        if feature is None:
            if near := override.get("coords_near"):
                row = Row(
                    "municipio",
                    m.code,
                    m.province,
                    existing_slugs.get(("municipio", m.code), ""),
                    {"es": m.parts[0]},
                    clean_aliases(m.parts, m.parts[0]),
                    population.get(m.code),
                    None,
                    None,
                )
                for lang in ("es", *NAME_LANGS):
                    if override.get(lang):
                        row.names[lang] = override[lang]
                rows.append(row)
                borrowed.append((row, m, near))
                continue
            report.without_geonames.append(f"{m.code} {m.raw_name}")
            continue
        if place:
            report.coords_from_place += 1
        else:
            report.coords_from_adm3 += 1

        # translations: town (place) names first, municipality (ADM3) names as fallback
        alt: dict[str, list[tuple[str, bool, bool]]] = defaultdict(list)
        for source in (place, area):
            if source:
                for lang, names in alt_names.get(source.id, {}).items():
                    alt[lang].extend(names)

        # Spanish display name: for bilingual INE names take the part GeoNames knows as Spanish
        es_known = {norm(n) for n, *_ in alt.get("es", [])}
        es_name = next((p for p in m.parts if norm(p) in es_known), m.parts[0])

        names = {"es": es_name}
        for lang in NAME_LANGS:
            if value := best_name(alt.get(lang, [])):
                names[lang] = value

        for lang in ("es", *NAME_LANGS):
            if override.get(lang):
                names[lang] = override[lang]

        alias_pool = [
            *m.parts,
            *override.get("aliases", []),
            *(n for lang in ALIAS_LANGS for n, *_ in alt.get(lang, [])),
            *(transliterate(names[lang], lang) for lang in ("uk", "ru") if lang in names),
        ]
        rows.append(
            Row(
                level="municipio",
                ine_code=m.code,
                parent_code=m.province,
                slug=existing_slugs.get(("municipio", m.code), ""),
                names=names,
                aliases=clean_aliases(alias_pool, es_name),
                population=population.get(m.code) or (feature.population or None),
                lon=feature.lon,
                lat=feature.lat,
                geonames_id=int(feature.id),
            )
        )

    for row, m, near in borrowed:
        source = next(
            (
                r
                for r in rows
                if r.parent_code == m.province and r.lat is not None and norm(r.names["es"]) == norm(near)
            ),
            None,
        )
        if source is None:
            rows.remove(row)
            report.without_geonames.append(f"{m.code} {m.raw_name} (coords_near '{near}' not found)")
            continue
        row.lat, row.lon = source.lat, source.lon
        report.approximate_coords.append(f"{m.code} {m.raw_name}: coordinates of {source.names['es']}")

    assign_municipality_slugs(rows, reserved_slugs, province_slugs, existing_slugs, report)
    return rows


def assign_municipality_slugs(
    rows: list[Row],
    reserved: set[str],
    province_slugs: dict[str, str],
    existing_slugs: dict[tuple[str, str], str],
    report: ImportReport,
) -> None:
    taken = set(existing_slugs.values()) | reserved
    new_rows = [r for r in rows if not r.slug]
    groups: dict[str, list[Row]] = defaultdict(list)
    for row in new_rows:
        groups[slugify(row.names["es"])].append(row)

    for base, group in groups.items():
        group.sort(key=lambda r: -(r.population or 0))
        for i, row in enumerate(group):
            candidate = base
            if i > 0 or candidate in taken:
                candidate = f"{base}-{province_slugs[row.parent_code]}"  # type: ignore[index]
                report.slug_collisions.append(f"{row.names['es']} ({row.ine_code}) -> {candidate}")
            if candidate in taken:
                candidate = f"{candidate}-{row.ine_code}"
            row.slug = candidate
            taken.add(candidate)


def build_localities(
    geo: GeoNames,
    municipalities: list[Row],
    alt_names: dict[str, dict[str, list[tuple[str, bool, bool]]]],
    existing_slugs: dict[tuple[str, str], str],
) -> list[Row]:
    """Villages, pedanías, parroquias: every GeoNames populated place of a municipality except its seat."""
    rows: list[Row] = []
    taken = set(existing_slugs.values())
    # a place used as a municipality's own point (also via overrides) is never a locality
    municipality_ids = {m.geonames_id for m in municipalities}
    for municipality in municipalities:
        seat_names = {norm(municipality.names["es"]), *(norm(a) for a in municipality.aliases)}
        for place in geo.places.get(municipality.ine_code or "", []):
            if int(place.id) in municipality_ids or norm(place.name) in seat_names:
                continue
            if NOT_A_LOCALITY.search(norm(place.name)):
                continue
            alt = alt_names.get(place.id, {})
            names = {"es": place.name}
            for lang in NAME_LANGS:
                if value := best_name(alt.get(lang, [])):
                    names[lang] = value
            aliases = [
                *(n for lang in ALIAS_LANGS for n, *_ in alt.get(lang, [])),
                *(transliterate(names[lang], lang) for lang in ("uk", "ru") if lang in names),
            ]

            slug = existing_slugs.get(("localidad", place.id))
            if not slug:
                slug = f"localidad-{slugify(place.name)}-{municipality.slug}"
                if slug in taken:
                    slug = f"{slug}-{place.id}"
                taken.add(slug)

            rows.append(
                Row(
                    level="localidad",
                    ine_code=None,
                    parent_code=municipality.ine_code,
                    slug=slug,
                    names=names,
                    aliases=clean_aliases(aliases, place.name),
                    population=place.population or None,
                    lon=place.lon,
                    lat=place.lat,
                    geonames_id=int(place.id),
                )
            )
    return rows


def build_regions(
    regions: dict[str, Any],
    register: list[IneMunicipality],
    municipalities: list[Row],
) -> tuple[list[Row], list[Row]]:
    province_to_community = {m.province: m.community for m in register}
    by_province: dict[str, list[Row]] = defaultdict(list)
    for row in municipalities:
        by_province[row.parent_code].append(row)  # type: ignore[index]

    def aggregate(rows: list[Row]) -> tuple[int, Row | None]:
        biggest = max(rows, key=lambda r: r.population or 0, default=None)
        return sum(r.population or 0 for r in rows), biggest

    provinces = []
    for p in regions["provincias"]:
        total, capital = aggregate(by_province[p["code"]])
        provinces.append(
            Row(
                level="provincia",
                ine_code=p["code"],
                parent_code=province_to_community.get(p["code"]),
                slug=f"provincia-{p['slug']}",
                names=p["names"],
                aliases=p.get("aliases", []),
                population=total or None,
                lon=capital.lon if capital else None,
                lat=capital.lat if capital else None,
            )
        )

    communities = []
    for c in regions["comunidades"]:
        rows = [
            r for code, rs in by_province.items() if province_to_community.get(code) == c["code"] for r in rs
        ]
        total, capital = aggregate(rows)
        communities.append(
            Row(
                level="comunidad",
                ine_code=c["code"],
                parent_code=None,
                slug=f"comunidad-{c['slug']}",
                names=c["names"],
                aliases=c.get("aliases", []),
                population=total or None,
                lon=capital.lon if capital else None,
                lat=capital.lat if capital else None,
            )
        )
    return communities, provinces


# ---------------------------------------------------------------------------- database


async def reserved_slugs(session: AsyncSession) -> set[str]:
    """Category and section slugs in every language: a location slug must never equal them."""
    reserved = {"oferta"}
    for model in (Category, Section):
        for (slug,) in await session.execute(select(model.slug)):
            reserved.update(slug.values())
    return reserved


async def upsert_rows(session: AsyncSession, rows: list[Row], parent_ids: dict[str, int]) -> dict[str, int]:
    if not rows:
        return {}
    level = rows[0].level
    values = [
        {
            "level": r.level,
            "ine_code": r.ine_code,
            "geonames_id": r.geonames_id,
            "parent_id": parent_ids.get(r.parent_code) if r.parent_code else None,
            "slug": r.slug,
            "names": r.names,
            "aliases": r.aliases,
            "population": r.population,
            "geog": WKTElement(f"POINT({r.lon} {r.lat})", srid=4326) if r.lon is not None else None,
        }
        for r in rows
    ]
    stmt = insert(Location)
    stmt = stmt.on_conflict_do_update(
        # localities have no INE code and are identified by their GeoNames id
        constraint="uq_locations_geonames_id" if level == "localidad" else "uq_locations_level_ine_code",
        set_={  # slug is intentionally not updated: URLs stay stable across re-imports
            "parent_id": stmt.excluded.parent_id,
            "geonames_id": stmt.excluded.geonames_id,
            "names": stmt.excluded.names,
            "aliases": stmt.excluded.aliases,
            "population": stmt.excluded.population,
            "geog": stmt.excluded.geog,
        },
    )
    for start in range(0, len(values), 1000):
        await session.execute(stmt, values[start : start + 1000])

    result = await session.execute(select(Location.ine_code, Location.id).where(Location.level == level))
    return dict(result.all())


async def import_locations(
    session: AsyncSession,
    raw_dir: Path = RAW_DIR,
    download: bool = False,
    log: Log = print,
) -> ImportReport:
    report = ImportReport()
    paths = ensure_sources(raw_dir, download, log)

    log("reading INE register and population")
    register = read_ine_register(paths["ine_register"])
    population = read_ine_population(paths["ine_population"])

    log("reading GeoNames ES dump")
    geo = read_geonames(paths["geonames"])
    ids = (
        {f.id for f in geo.adm3.values()}
        | {f.id for fs in geo.places.values() for f in fs}
        | {f.id for fs in geo.places_by_name.values() for f in fs}
    )
    alt_names = read_alternate_names(paths["geonames_alt"], ids, log)

    regions = json.loads(REGIONS_FILE.read_text(encoding="utf-8"))
    overrides = json.loads(OVERRIDES_FILE.read_text(encoding="utf-8")).get("municipios", {})
    province_slugs = {p["code"]: p["slug"] for p in regions["provincias"]}

    existing = await session.execute(
        select(Location.level, Location.ine_code, Location.geonames_id, Location.slug)
    )
    existing_slugs = {
        (level, str(gid) if level == "localidad" else code): slug for level, code, gid, slug in existing
    }

    municipalities = build_municipalities(
        register,
        population,
        geo,
        alt_names,
        overrides,
        await reserved_slugs(session),
        province_slugs,
        existing_slugs,
        report,
    )
    communities, provinces = build_regions(regions, register, municipalities)
    localities = build_localities(geo, municipalities, alt_names, existing_slugs)

    log("writing to the database")
    community_ids = await upsert_rows(session, communities, {})
    province_ids = await upsert_rows(session, provinces, community_ids)
    municipality_ids = await upsert_rows(session, municipalities, province_ids)
    await upsert_rows(session, localities, municipality_ids)
    # localities are fully derived from GeoNames: drop the ones that disappeared or got filtered out
    await session.execute(
        delete(Location).where(
            Location.level == "localidad",
            Location.geonames_id.not_in([r.geonames_id for r in localities]),
        )
    )
    await session.commit()

    report.comunidades, report.provincias, report.municipios, report.localidades = (
        len(communities),
        len(provinces),
        len(municipalities),
        len(localities),
    )
    report.top = sorted(municipalities, key=lambda r: -(r.population or 0))[:60]
    return report
